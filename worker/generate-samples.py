# Run through uv with python-docx, python-pptx, odfpy and Pillow.
import sys
from pathlib import Path
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from odf.opendocument import OpenDocumentText, OpenDocumentPresentation
from odf.text import P
from odf.draw import Page, Frame, TextBox
from odf.style import Style, MasterPage, PageLayout
from PIL import Image, ImageDraw
p = Path(sys.argv[1])
im = Image.new('RGB', (500, 150), '#148a65'); ImageDraw.Draw(im).text((20, 30), 'Embedded acceptance image', fill='white'); im.save(p/'image.png')
d = Document(); d.add_heading('DOCX acceptance sample', 0); d.add_paragraph('Paragraph with bold content.').add_run(' Bold.').bold = True; d.add_picture(str(p/'image.png'), width=Inches(3)); d.save(p/'sample.docx')
prs = Presentation(); prs.slide_width = Inches(13.333333); prs.slide_height = Inches(7.5)
for i in range(2):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor(235, 245, 255)
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(1)); run = box.text_frame.paragraphs[0].add_run(); run.text = f'PPTX acceptance sample slide {i+1}'; run.font.size = Pt(32); run.font.bold = True
    slide.shapes.add_picture(str(p/'image.png'), Inches(1), Inches(3), width=Inches(4))
prs.save(p/'sample.pptx')
odt = OpenDocumentText(); odt.text.addElement(P(text='ODT acceptance sample')); odt.save(str(p/'sample.odt'))
odp = OpenDocumentPresentation(); style = Style(name='Default', family='drawing-page'); odp.automaticstyles.addElement(style); odp.automaticstyles.addElement(PageLayout(name='layout')); odp.masterstyles.addElement(MasterPage(name='Default', pagelayoutname='layout'))
for i in range(2):
    page = Page(name=f'page{i+1}', stylename=style, masterpagename='Default'); frame = Frame(x='2cm', y='2cm', width='20cm', height='3cm'); box = TextBox(); box.addElement(P(text=f'ODP acceptance sample slide {i+1}')); frame.addElement(box); page.addElement(frame); odp.presentation.addElement(page)
odp.save(str(p/'sample.odp'))
